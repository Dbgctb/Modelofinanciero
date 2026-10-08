import openpyxl
from collections import defaultdict

wb = openpyxl.load_workbook('diario 2025 david.xlsx', read_only=True, data_only=True)

# Hoja Recibos: todos los codigos y sus entes
ws_r = wb['Recibos']
recibos_hoja = {}        # codigo -> ente
recibos_hoja_ente = defaultdict(set)
total_filas_recibos = 0
for i, row in enumerate(ws_r.iter_rows(values_only=True)):
    if i == 0:
        continue
    if row[0] is None:
        break
    total_filas_recibos += 1
    try:
        cod = str(int(row[0]))
    except (TypeError, ValueError):
        cod = str(row[0]).strip()
    ente = str(row[10]).strip() if row[10] else '(sin ente)'
    recibos_hoja[cod] = ente
    recibos_hoja_ente[ente].add(cod)

print('=== HOJA RECIBOS ===')
print('Filas totales:', total_filas_recibos)
print('Codigos unicos:', len(recibos_hoja))
print()

# Hoja completada: recibos referenciados en cuentas 6 y 7
ws_d = wb['Hoja completada']
recibos_diario = {}      # recibo_str -> {cia, aux, agrup}
recibos_diario_aux = defaultdict(set)
recibos_diario_cia = defaultdict(set)

for i, row in enumerate(ws_d.iter_rows(values_only=True)):
    if i == 0:
        continue
    if row[0] is None:
        break
    agrup = str(row[9]) if row[9] else ''
    recibo = row[12]
    if not recibo or str(recibo).strip() in ('', 'None', '0'):
        continue
    r = str(recibo).strip()
    if r not in recibos_diario:
        recibos_diario[r] = {
            'cia': str(row[13]).strip() if row[13] else '',
            'aux': str(row[14]).strip() if row[14] else '',
            'agrup': agrup,
        }
    cia = str(row[13]).strip() if row[13] else '(sin cia)'
    aux = str(row[14]).strip() if row[14] else '(sin aux)'
    recibos_diario_cia[cia].add(r)
    recibos_diario_aux[aux].add(r)

print('=== DIARIO (ctas 6+7) ===')
print('Recibos unicos referenciados:', len(recibos_diario))
print()

# Cruce
solo_diario = set(recibos_diario.keys()) - set(recibos_hoja.keys())
solo_hoja   = set(recibos_hoja.keys()) - set(recibos_diario.keys())
en_ambos    = set(recibos_diario.keys()) & set(recibos_hoja.keys())

print('=== CRUCE ===')
print('En ambos (OK):', len(en_ambos))
print('En diario PERO NO en hoja Recibos:', len(solo_diario))
print('En hoja Recibos PERO NO en diario:', len(solo_hoja))

if solo_diario:
    print()
    print('Recibos del diario SIN ficha en hoja Recibos (todos):')
    for r in sorted(solo_diario, key=lambda x: int(x) if x.isdigit() else 0):
        info = recibos_diario[r]
        cia = info['cia'][:35]
        aux = info['aux'][:35]
        cta = info['agrup']
        print('  Recibo', r, ' cta=' + cta, ' cia=' + cia, ' aux=' + aux)

# Recibos por auxiliar: comparacion diario vs hoja recibos
print()
print('=== RECIBOS POR AUXILIAR (diario vs hoja) ===')
print('Auxiliar                                       Diario  HojaRec')
todos_aux = sorted(set(list(recibos_diario_aux.keys()) + list(recibos_hoja_ente.keys())))
for aux in todos_aux:
    n_d = len(recibos_diario_aux.get(aux, set()))
    n_h = len(recibos_hoja_ente.get(aux, set()))
    diff = n_d - n_h
    flag = ' <-- DIFIERE' if diff != 0 else ''
    if n_d > 0 or n_h > 0:
        print(aux[:45].ljust(46), str(n_d).rjust(6), str(n_h).rjust(8), flag)

# Recibos por compania (diario)
print()
print('=== RECIBOS POR COMPANIA (diario) ===')
recibos_diario_cia_705 = defaultdict(set)
recibos_diario_cia_623 = defaultdict(set)
for i, row in enumerate(ws_d.iter_rows(values_only=True)):
    if i == 0:
        continue
    if row[0] is None:
        break
    agrup = str(row[9]) if row[9] else ''
    recibo = row[12]
    if not recibo or str(recibo).strip() in ('', 'None', '0'):
        continue
    r = str(recibo).strip()
    cia = str(row[13]).strip() if row[13] else '(sin cia)'
    if agrup.startswith('705'):
        recibos_diario_cia_705[cia].add(r)
    elif agrup.startswith('623'):
        recibos_diario_cia_623[cia].add(r)

all_cias = sorted(set(list(recibos_diario_cia_705.keys()) + list(recibos_diario_cia_623.keys())))
print('Compania                           705-recibos  623-recibos')
for cia in all_cias:
    n705 = len(recibos_diario_cia_705.get(cia, set()))
    n623 = len(recibos_diario_cia_623.get(cia, set()))
    print(cia[:34].ljust(35), str(n705).rjust(11), str(n623).rjust(12))
